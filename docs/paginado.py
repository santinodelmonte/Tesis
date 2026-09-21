# -*- coding: utf-8 -*-
"""Numera las paginas de un .docx generado.

Los documentos vienen de una herramienta que dejo el pie de pagina sin enganchar a
ninguna seccion y la numeracion arrancando en cero. El resultado es un trabajo de mas
de cien paginas sin un solo numero impreso, y un indice que no puede referirse a nada.
Se arregla en el momento de generar, que es donde se arregla todo en este repositorio.

Lo que hace, por seccion:

* saca el `w:pgNumType w:start="0"`, que es de donde salia el numero cero;
* engancha un pie con el campo PAGE centrado;
* en la primera seccion marca la primera pagina como distinta y le deja el pie vacio,
  para que la portada no lleve numero. En las demas no: la primera pagina de una
  seccion interna es una pagina cualquiera y tiene que llevarlo.
"""
from docx.oxml.ns import qn
from docx.shared import Pt

CAMPO = 'PAGE'


def _campo_pagina(parrafo):
    """Escribe el campo PAGE como lo espera Word: begin, instruccion, separate, end."""
    from docx.oxml import OxmlElement

    def marca(tipo):
        run = OxmlElement('w:r')
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'), tipo)
        run.append(fld)
        return run

    run = OxmlElement('w:r')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' %s ' % CAMPO
    run.append(instr)

    p = parrafo._p
    p.append(marca('begin'))
    p.append(run)
    p.append(marca('separate'))
    # El resultado en cache: Word lo reemplaza al abrir, y LibreOffice al exportar.
    resultado = OxmlElement('w:r')
    texto = OxmlElement('w:t')
    texto.text = '1'
    resultado.append(texto)
    p.append(resultado)
    p.append(marca('end'))


def _vaciar(contenedor):
    for p in list(contenedor.paragraphs):
        p._p.getparent().remove(p._p)


def numerar(documento, tamanio=Pt(10)):
    """Deja el documento con numero de pagina en todas las paginas menos la portada."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    secciones = 0
    for indice, seccion in enumerate(documento.sections):
        sect = seccion._sectPr
        for nodo in sect.findall(qn('w:pgNumType')):
            sect.remove(nodo)

        # Solo la portada va sin numero, y la portada esta en la primera seccion.
        portada_aparte = indice == 0
        seccion.different_first_page_header_footer = portada_aparte

        pie = seccion.footer
        pie.is_linked_to_previous = False
        _vaciar(pie)
        parrafo = pie.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _campo_pagina(parrafo)
        for run in parrafo.runs:
            run.font.size = tamanio

        if portada_aparte:
            portada = seccion.first_page_footer
            portada.is_linked_to_previous = False
            _vaciar(portada)
            portada.add_paragraph()

        secciones += 1
    return secciones
